import io
import re
from datetime import datetime
from pathlib import Path
import pandas as pd
from openpyxl.styles import PatternFill
from ..catalog.schema import CanonicalSchema, load_schema
from ..common.models import CanonicalRow, FileResult

_TRACE_COLUMNS = [
    ("source_file", "Arquivo de origem"),
    ("source_sheet", "Aba de origem"),
    ("format_id", "Modelo"),
    ("confidence", "Confiança"),
    ("needs_review", "Revisar"),
    ("ocr_used", "OCR"),
    ("row_index_in_source", "Linha na origem"),
]


def assemble_output(
    canonical_rows: list[CanonicalRow],
    file_results: list[FileResult],
    output_path: Path | str | None = None,
    schema: CanonicalSchema | None = None,
) -> bytes:
    """Assemble compiled output xlsx. Returns bytes (for Streamlit download).

    A aba "Dados" tem as colunas padrão primeiro (na ordem do esquema, só as que
    algum modelo preencheu), seguidas das colunas de rastreabilidade.
    """
    schema = schema or load_schema()
    df_data = _build_data_frame(canonical_rows, schema)

    log_records = [
        {
            "arquivo": Path(r.path).name,
            "formato": r.format_id or "—",
            "confianca": round(r.confidence, 3),
            "status": r.status,
            "linhas_extraidas": r.rows_extracted,
            "erro": r.error or "",
        }
        for r in file_results
    ]
    df_log = pd.DataFrame(log_records)

    output = io.BytesIO()
    with pd.ExcelWriter(
        output, engine="openpyxl", date_format="DD/MM/YYYY", datetime_format="DD/MM/YYYY"
    ) as writer:
        df_data.to_excel(writer, sheet_name="Dados", index=False)
        df_log.to_excel(writer, sheet_name="Log", index=False)
        _style_log_sheet(writer.book["Log"], file_results)

    output.seek(0)
    result = output.read()
    if output_path:
        Path(output_path).write_bytes(result)
    return result


def _build_data_frame(canonical_rows: list[CanonicalRow], schema: CanonicalSchema) -> pd.DataFrame:
    if not canonical_rows:
        return pd.DataFrame()
    used = {k for r in canonical_rows for k in r.data}
    ordered = [f.name for f in schema.fields if f.name in used]
    ordered += sorted(used - set(ordered))  # campos de modelos antigos fora do esquema

    records = []
    for r in canonical_rows:
        rec = {}
        for name in ordered:
            rec[schema.label(name)] = _coerce(r.data.get(name), schema.get(name))
        rec.update({
            "source_file": Path(r.source_file).name,
            "source_sheet": r.source_sheet,
            "format_id": r.format_id or "auto (sem modelo cadastrado)",
            "confidence": round(r.confidence, 3),
            "needs_review": r.needs_review,
            "ocr_used": r.ocr_used,
            "row_index_in_source": r.row_index_in_source,
        })
        records.append(rec)
    return pd.DataFrame(records).rename(columns=dict(_TRACE_COLUMNS))


def _coerce(value, field):
    """Converte para número/data conforme o tipo da coluna padrão; se não der, mantém o texto."""
    if value is None or field is None or field.type == "text":
        return value
    text = str(value).strip()
    if not text:
        return None
    if field.type == "number":
        return _to_number(text)
    return _to_date(text)


def _to_number(text: str):
    s = text.replace("R$", "").replace(" ", "")
    if "," in s:  # formato brasileiro: 1.234,56
        s = s.replace(".", "").replace(",", ".")
    try:
        return float(s)
    except ValueError:
        return text


def _to_date(text: str):
    iso = re.match(r"^\d{4}-\d{2}-\d{2}", text)
    try:
        if iso:
            return datetime.fromisoformat(text[:19])
        return pd.to_datetime(text, dayfirst=True, format="mixed").to_pydatetime()
    except (ValueError, TypeError):
        return text


def _style_log_sheet(ws, file_results: list[FileResult]) -> None:
    GREEN = PatternFill("solid", fgColor="C6EFCE")
    YELLOW = PatternFill("solid", fgColor="FFEB9C")
    RED = PatternFill("solid", fgColor="FFC7CE")
    for i, result in enumerate(file_results, start=2):
        fill = GREEN if result.status == "ok" else YELLOW if result.status == "review" else RED
        for cell in ws[i]:
            cell.fill = fill
